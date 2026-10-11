from bhb.gold import GoldRow, score_records, write_sample
from bhb.model import Host, Person, Record


def _record() -> Record:
    return Record(
        id="bhb:1992-59457",
        source="print-straive",
        bgb_year=1992,
        print_number="59457",
        type="book",
        title="Images du scoutisme : 50 ans de calendrier FSC",
        persons=[Person(family="Steeman", given="Stéphane")],
        year=1991,
        publisher=None,
        host=Host(title="Some host"),
    )


def test_score_matches_a_normalised_title_and_surname():
    gold = [
        GoldRow(
            label="59457",
            bgb_year=1992,
            title="Images du scoutisme : 50 ans de calendrier FSC.",
            host_title="",
            surname="Steeman",
            year="1991",
            kind="book",
        )
    ]
    score = score_records([_record()], gold)
    assert score.rate("title") == 1
    assert score.rate("surname") == 1
    assert score.rate("year") == 1
    assert score.compared["host"] == 0
    assert score.missing_records == 0


def test_the_full_sample_adds_pre1993_rows_and_a_theme_column(tmp_path):
    records = []
    for index in range(180):
        records.append(
            Record(id=f"bhb:1960-{index}", source="print-straive", bgb_year=1960, print_number=str(index), role="entry", title="t")
        )
    for index in range(40):
        records.append(
            Record(id=f"bhb:2000-{index}", source="print-straive", bgb_year=2000, print_number=str(index), role="entry", title="t")
        )
    path = tmp_path / "sample.xlsx"
    assert write_sample(records, path, n=300, seed=1) == 220
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True)
    rows = list(book.active.iter_rows(values_only=True))
    book.close()
    assert rows[0][-1] == "theme"
    ids = [row[0] for row in rows[1:]]
    assert len(ids) == len(set(ids)) == 220


def test_sample_spreadsheet_round_trip(tmp_path):
    path = tmp_path / "sample.xlsx"
    assert write_sample([_record()], path, n=10, seed=1) == 1
    from bhb.gold import load_gold

    rows = load_gold([path])
    assert rows[0].label == "59457"
    assert rows[0].bgb_year == 1992
    assert rows[0].surname == "Steeman"
